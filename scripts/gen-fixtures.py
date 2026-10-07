"""Generates every file in public/fixtures. Pure Python + openpyxl + ffmpeg (for y4m/mjpeg/jpeg).
All people, numbers and documents are fictional."""
import csv, io, json, os, struct, subprocess, zipfile, zlib
OUT = 'public/fixtures'
os.makedirs(f'{OUT}/testdata', exist_ok=True); os.makedirs(f'{OUT}/extension', exist_ok=True)

# ---------- tiny 5x7 bitmap font ----------
F = {
'A':"01110 10001 10001 11111 10001 10001 10001",'B':"11110 10001 10001 11110 10001 10001 11110",
'C':"01111 10000 10000 10000 10000 10000 01111",'D':"11110 10001 10001 10001 10001 10001 11110",
'E':"11111 10000 10000 11110 10000 10000 11111",'F':"11111 10000 10000 11110 10000 10000 10000",
'G':"01111 10000 10000 10111 10001 10001 01111",'H':"10001 10001 10001 11111 10001 10001 10001",
'I':"11111 00100 00100 00100 00100 00100 11111",'J':"00111 00010 00010 00010 10010 10010 01100",
'K':"10001 10010 10100 11000 10100 10010 10001",'L':"10000 10000 10000 10000 10000 10000 11111",
'M':"10001 11011 10101 10101 10001 10001 10001",'N':"10001 11001 10101 10011 10001 10001 10001",
'O':"01110 10001 10001 10001 10001 10001 01110",'P':"11110 10001 10001 11110 10000 10000 10000",
'Q':"01110 10001 10001 10001 10101 10010 01101",'R':"11110 10001 10001 11110 10100 10010 10001",
'S':"01111 10000 10000 01110 00001 00001 11110",'T':"11111 00100 00100 00100 00100 00100 00100",
'U':"10001 10001 10001 10001 10001 10001 01110",'V':"10001 10001 10001 10001 10001 01010 00100",
'W':"10001 10001 10001 10101 10101 10101 01010",'X':"10001 10001 01010 00100 01010 10001 10001",
'Y':"10001 10001 01010 00100 00100 00100 00100",'Z':"11111 00001 00010 00100 01000 10000 11111",
'0':"01110 10001 10011 10101 11001 10001 01110",'1':"00100 01100 00100 00100 00100 00100 01110",
'2':"01110 10001 00001 00010 00100 01000 11111",'3':"11110 00001 00001 01110 00001 00001 11110",
'4':"00010 00110 01010 10010 11111 00010 00010",'5':"11111 10000 11110 00001 00001 10001 01110",
'6':"00110 01000 10000 11110 10001 10001 01110",'7':"11111 00001 00010 00100 01000 01000 01000",
'8':"01110 10001 10001 01110 10001 10001 01110",'9':"01110 10001 10001 01111 00001 00010 01100",
'-':"00000 00000 00000 11111 00000 00000 00000",':':"00000 00100 00000 00000 00000 00100 00000",
'/':"00001 00010 00010 00100 01000 01000 10000",'.':"00000 00000 00000 00000 00000 00000 00100",
' ':"00000 00000 00000 00000 00000 00000 00000",
}
def png(path, w, h, bg, texts, rects=()):
    px = [[bg]*w for _ in range(h)]
    for (x0,y0,x1,y1,c) in rects:
        for y in range(max(0,y0),min(h,y1)):
            for x in range(max(0,x0),min(w,x1)): px[y][x]=c
    for (text,x,y,scale,c) in texts:
        cx=x
        for ch in text.upper():
            rows=F.get(ch,F[' ']).split()
            for ry,row in enumerate(rows):
                for rx,bit in enumerate(row):
                    if bit=='1':
                        for dy in range(scale):
                            for dx in range(scale):
                                yy,xx=y+ry*scale+dy,cx+rx*scale+dx
                                if 0<=yy<h and 0<=xx<w: px[yy][xx]=c
            cx+=6*scale
    raw=b''.join(b'\x00'+b''.join(bytes(p) for p in row) for row in px)
    def chunk(t,d): return struct.pack('>I',len(d))+t+d+struct.pack('>I',zlib.crc32(t+d)&0xffffffff)
    data=b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',w,h,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(raw,9))+chunk(b'IEND',b'')
    open(path,'wb').write(data)

WHITE=(255,255,255); DARK=(20,30,45)
def id_card(side, bg):
    texts=[('SAMPLE IDENTITY DOCUMENT',40,30,3,WHITE),(side,40,90,10,WHITE)]
    rects=[(380,170,600,370,(255,255,255))] if side=='FRONT' else [(40,300,600,330,(10,10,10)),(40,340,600,360,(10,10,10))]
    if side=='FRONT':
        texts+=[('NAME: ALEX SAMPLE',40,200,3,WHITE),('NO: 000-TEST-0001',40,240,3,WHITE),('DOB: 01/01/1990',40,280,3,WHITE),('PHOTO',440,260,3,DARK)]
    else:
        texts+=[('ISSUED BY: TEST AUTHORITY',40,200,3,WHITE),('VALID UNTIL: 12/2030',40,240,3,WHITE)]
    png(f'{OUT}/{side.lower()}.png',640,400,bg,texts,rects)
id_card('FRONT',(31,157,85))    # green
id_card('BACK',(37,99,235))     # blue
png(f'{OUT}/avatar.png',256,256,(124,58,237),[('AS',58,90,12,WHITE)])
png(f'{OUT}/receipt.png',400,520,(250,250,250),[('RECEIPT',40,30,5,DARK),('BLUE MUG X2',40,120,3,DARK),('TOTAL 24.00',40,170,3,DARK)])

def ff(*a): subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y',*a],check=True)
for side in ('front','back'):
    ff('-i',f'{OUT}/{side}.png','-q:v','3',f'{OUT}/{side}.jpg')
    ff('-loop','1','-i',f'{OUT}/{side}.png','-frames:v','5','-r','5','-pix_fmt','yuv420p',f'{OUT}/{side}.y4m')
    ff('-loop','1','-i',f'{OUT}/{side}.png','-frames:v','5','-r','5','-q:v','3','-f','mjpeg',f'{OUT}/{side}.mjpeg')
# 5 MB heavy image (noise so it doesn't compress)
ff('-f','lavfi','-i','nullsrc=s=1600x1200,geq=random(1)*255:128:128','-frames:v','1',f'{OUT}/heavy.png')

# ---------- CSV ----------
people=[('Ada Lovelace','ada@example.com','admin'),('Grace Hopper','grace@example.com','approver'),('Alan Turing','alan@example.com','viewer'),
 ('Katherine Johnson','katherine@example.com','viewer'),('Linus Example','linus@example.com','approver'),('Margaret Hamilton','margaret@example.com','admin'),
 ('Tim Sample','tim@example.com','viewer'),('Barbara Liskov','barbara@example.com','approver'),('Dennis Test','dennis@example.com','viewer'),('Radia Perlman','radia@example.com','viewer')]
with open(f'{OUT}/sample.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['name','email','role']); w.writerows(people)
with open(f'{OUT}/bad-missing-email.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['name','role']); w.writerows([(p[0],p[2]) for p in people])
with open(f'{OUT}/bad-empty.csv','w') as f: f.write('')
open(f'{OUT}/notes.txt','w').write('Plain text fixture.\nLine two.\n')

# test-data profiles (data-driven)
reg=[('valid-1','Ada','Lovelace','ada.l@example.com','Str0ng!Pass','Str0ng!Pass','+44 20 7946 0001','true','Account created'),
     ('valid-2','Grace','Hopper','grace.h@example.com','Navy#1906pw','Navy#1906pw','+1 202 555 0143','true','Account created'),
     ('bad-email','Alan','Turing','alan.example.com','Enigma!1936','Enigma!1936','+44 20 7946 0002','true','Enter a valid email address'),
     ('weak-pass','Katherine','Johnson','kj@example.com','abc','abc','+1 202 555 0199','true','Password must be at least 8 characters'),
     ('mismatch','Radia','Perlman','rp@example.com','Spanning#Tree1','Spanning#Tree2','+1 202 555 0111','true','Passwords do not match'),
     ('no-terms','Barbara','Liskov','bl@example.com','Substitut!on1','Substitut!on1','+1 202 555 0177','false','You must accept the terms')]
hdr=['case','firstName','lastName','email','password','confirmPassword','phone','acceptTerms','expectedMessage']
with open(f'{OUT}/testdata/registration.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(hdr); w.writerows(reg)
import openpyxl
wb=openpyxl.Workbook(); ws=wb.active; ws.title='registration'; ws.append(hdr); [ws.append(list(r)) for r in reg]; wb.save(f'{OUT}/testdata/registration.xlsx')
survey=[('happy',5,'9','Yes','Great product'),('neutral',3,'6','No','Okay'),('unhappy',1,'2','No','Too slow')]
with open(f'{OUT}/testdata/survey.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['case','stars','nps','recommend','comment']); w.writerows(survey)

# line items xlsx (+ known totals)
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Line items'
ws.append(['sku','description','quantity','unit_price'])
items=[('SKU-100','Laptop stand',2,45.00),('SKU-200','USB-C hub',3,39.50),('SKU-300','Desk lamp',1,64.00),('SKU-400','Monitor arm',2,89.00),('SKU-500','Cable kit',5,12.00)]
for i in items: ws.append(list(i))
wb.save(f'{OUT}/line-items.xlsx')
wb=openpyxl.Workbook(); ws=wb.active; ws.append(['sku','qty']); ws.append(['SKU-1',1]); wb.save(f'{OUT}/line-items-bad-headers.xlsx')

# ---------- PDF (hand-written, text only) ----------
def pdf(path, lines):
    content='BT /F1 12 Tf 50 780 Td 16 TL '+' '.join(f'({l.replace("(","[").replace(")","]")}) Tj T*' for l in lines)+' ET'
    objs=['<< /Type /Catalog /Pages 2 0 R >>','<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
          '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
          f'<< /Length {len(content)} >>\nstream\n{content}\nendstream','<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    out=b'%PDF-1.4\n'; offs=[]
    for i,o in enumerate(objs,1):
        offs.append(len(out)); out+=f'{i} 0 obj\n{o}\nendobj\n'.encode()
    x=len(out); out+=f'xref\n0 {len(objs)+1}\n0000000000 65535 f \n'.encode()+b''.join(f'{o:010d} 00000 n \n'.encode() for o in offs)
    out+=f'trailer\n<< /Size {len(objs)+1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n'.encode()
    open(path,'wb').write(out)
inv=['INVOICE INV-2026-0042','Bill to: Example Corp, 1 Sample Street','Date: 06 October 2026','',
     'Line items:']+[f'{s}  {d}  qty {q}  @ {p:.2f}  = {q*p:.2f}' for s,d,q,p in items]+['','Total: $1,234.50'.replace('1,234.50',f'{sum(q*p for _,_,q,p in items):,.2f}'),'Status: Paid']
pdf(f'{OUT}/invoice.pdf',inv)
pdf(f'{OUT}/terms.pdf',['Terms of Use (sample)','1. This is a fictional document for testing.','2. No real agreement is formed.'])

# ---------- test browser extension ----------
manifest={"manifest_version":3,"name":"Test Playground Extension","version":"1.0.0",
 "description":"Adds a visible banner to playground pages so tests can verify the extension loaded.",
 "content_scripts":[{"matches":["<all_urls>"],"js":["content.js"],"run_at":"document_idle"}]}
content="""(function(){if(document.getElementById('test-extension-banner'))return;var b=document.createElement('div');b.id='test-extension-banner';b.setAttribute('data-testid','extension-banner');b.textContent='Test extension active';b.style.cssText='position:fixed;top:0;left:50%;transform:translateX(-50%);z-index:2147483647;background:#0f766e;color:#fff;padding:4px 12px;font:600 12px system-ui;border-radius:0 0 8px 8px';document.documentElement.appendChild(b);document.documentElement.setAttribute('data-tp-extension','active');})();"""
with zipfile.ZipFile(f'{OUT}/extension/test-extension.zip','w',zipfile.ZIP_DEFLATED) as z:
    z.writestr('manifest.json',json.dumps(manifest,indent=2)); z.writestr('content.js',content)
total=sum(q*p for _,_,q,p in items)
json.dump({'lineItemsTotal':round(total,2),'lineItemsCount':len(items),'sampleCsvRows':len(people)},open(f'{OUT}/expected.json','w'),indent=1)
print('invoice total', total)
